import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from jinja2 import Environment, FileSystemLoader

from .config import BASE_DIR, DATA_DIR
from .db import connect, last_snapshot, new_notes_this_run

TEMPLATE_DIR = BASE_DIR / "monitor" / "templates"

EXCEL_HEADERS = ["账号", "标题", "类型", "发布时间", "点赞", "收藏", "评论", "分享", "话题", "IP属地", "链接"]


def _fmt_ts(ms):
    if not ms:
        return ""
    return datetime.fromtimestamp(ms / 1000).strftime("%Y-%m-%d %H:%M")


def build_user_summary(conn):
    rows = conn.execute(
        """SELECT user_id, MAX(nickname) AS nickname, COUNT(*) AS total,
             SUM(CASE WHEN first_seen_at >= date('now','localtime','start of day') THEN 1 ELSE 0 END) AS today_new,
             MAX(publish_time) AS last_publish
           FROM notes GROUP BY user_id"""
    ).fetchall()
    return [dict(r) for r in rows]


def build_hot_notes(conn, top=10):
    rows = conn.execute(
        "SELECT * FROM notes WHERE detail_fetched=1 ORDER BY publish_time DESC LIMIT 300"
    ).fetchall()
    growth = []
    for r in rows:
        prev = last_snapshot(conn, r["note_id"])
        if prev and r["liked_count"] > prev["liked_count"]:
            growth.append({
                "title": r["title"], "note_url": r["note_url"], "nickname": r["nickname"],
                "liked_growth": r["liked_count"] - prev["liked_count"],
                "liked_count": r["liked_count"],
                "collected_count": r["collected_count"], "comment_count": r["comment_count"],
            })
    growth.sort(key=lambda x: -x["liked_growth"])
    return growth[:top]


def build_excel(result, conn) -> Path:
    run_at = result["run_at"].replace(":", "-").replace(" ", "_")
    path = DATA_DIR / "excel" / f"竞品监测_{run_at}.xlsx"
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "账号汇总"
    head_fill = PatternFill("solid", fgColor="FFD700")
    for col, h in enumerate(["账号", "昵称", "笔记总数", "今日新增", "最新发布时间"], 1):
        c = ws1.cell(row=1, column=col, value=h)
        c.font = Font(bold=True)
        c.fill = head_fill
    for row, s in enumerate(result["stats"], 2):
        total_row = conn.execute(
            "SELECT COUNT(*), MAX(publish_time) FROM notes WHERE user_id=?", (s["user_id"],)
        ).fetchone()
        ws1.append([s["name"], s.get("nickname") or "", total_row[0], s["new_notes"],
                    _fmt_ts(total_row[1])])

    ws2 = wb.create_sheet("本期新增笔记")
    for col, h in enumerate(EXCEL_HEADERS, 1):
        c = ws2.cell(row=1, column=col, value=h)
        c.font = Font(bold=True)
        c.fill = head_fill
    for n in result.get("new_notes", []):
        ws2.append([n["user_id"], n["title"], n["note_type"], _fmt_ts(n["publish_time"]),
                    n["liked_count"], n["collected_count"], n["comment_count"],
                    n["share_count"], n.get("tags"), n.get("ip_location"), n["note_url"]])

    for ws in (ws1, ws2):
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(width + 4, 60)

    wb.save(path)
    return path


def build_html(result, conn) -> Path:
    env = Environment(loader=FileSystemLoader(TEMPLATE_DIR))
    tpl = env.get_template("report.html.j2")
    html = tpl.render(
        run_at=result["run_at"],
        duration=result["duration_s"],
        stats=result["stats"],
        user_summary=build_user_summary(conn),
        new_notes=result.get("new_notes", []),
        hot_notes=result.get("hot_notes", []),
        errors=result.get("errors", []),
        fmt_ts=_fmt_ts,
    )
    path = DATA_DIR / "reports" / f"report_{result['run_at'].replace(':', '-').replace(' ', '_')}.html"
    path.write_text(html, encoding="utf-8")
    return path
