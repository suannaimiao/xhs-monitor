"""周报：汇总近 7 天发布的竞品笔记，生成 HTML + Excel 并邮件发送。

用法：uv run python -m monitor.weekly [天数，默认 7]
"""
import sys
from datetime import datetime, timedelta
from pathlib import Path

from jinja2 import Environment, FileSystemLoader
from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .config import BASE_DIR, DATA_DIR
from .db import connect
from .mailer import send_report

TEMPLATE_DIR = BASE_DIR / "monitor" / "templates"
HEADERS = ["账号", "标题", "类型", "发布时间", "点赞", "收藏", "评论", "分享", "联名", "话题", "链接"]


def _fmt_ts(ms):
    if not ms:
        return ""
    return datetime.fromtimestamp(ms / 1000).strftime("%Y-%m-%d %H:%M")


def fetch_week(conn, days: int):
    since_ms = int((datetime.now() - timedelta(days=days)).timestamp() * 1000)
    notes = [dict(r) for r in conn.execute(
        """SELECT * FROM notes WHERE publish_time >= ? ORDER BY publish_time DESC""",
        (since_ms,)).fetchall()]
    accounts = [dict(r) for r in conn.execute(
        """SELECT user_id, MAX(nickname) AS nickname, COUNT(*) AS total,
             SUM(liked_count) AS likes, SUM(collected_count) AS collects,
             SUM(CASE WHEN cobrand != '' THEN 1 ELSE 0 END) AS cobrand_cnt
           FROM notes WHERE publish_time >= ? GROUP BY user_id ORDER BY total DESC""",
        (since_ms,)).fetchall()]
    return notes, accounts


def build_excel(notes, path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "本周笔记"
    for col, h in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFD700")
    for n in notes:
        ws.append([n["nickname"], n["title"], n["note_type"], _fmt_ts(n["publish_time"]),
                   n["liked_count"], n["collected_count"], n["comment_count"],
                   n["share_count"], n["cobrand"], n["tags"], n["note_url"]])
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(width + 4, 60)
    wb.save(path)
    return path


def main():
    days = int(sys.argv[1]) if len(sys.argv) > 1 else 7
    conn = connect()
    notes, accounts = fetch_week(conn, days)
    top = sorted(notes, key=lambda n: -n["liked_count"])[:10]
    cobrand_notes = [n for n in notes if n["cobrand"]]
    conn.close()

    since = (datetime.now() - timedelta(days=days)).strftime("%m-%d")
    until = datetime.now().strftime("%m-%d")

    env = Environment(loader=FileSystemLoader(TEMPLATE_DIR))
    html = env.get_template("weekly.html.j2").render(
        since=since, until=until, days=days,
        accounts=accounts, notes=notes, top=top, cobrand_notes=cobrand_notes,
        fmt_ts=_fmt_ts, total_likes=sum(n["liked_count"] for n in notes),
        cobrand_total=len(cobrand_notes),
    )

    stamp = f"{since.replace('-', '')}-{until.replace('-', '')}"
    excel_path = build_excel(notes, DATA_DIR / "excel" / f"竞品周报_{stamp}.xlsx")
    html_path = DATA_DIR / "reports" / f"weekly_{stamp}.html"
    html_path.write_text(html, encoding="utf-8")

    from .config import load_config
    cfg = load_config()
    subject = f"【XHS竞品监测】周报 {since}~{until}：{len(notes)} 篇笔记，{len(cobrand_notes)} 篇联名"
    ok = send_report(cfg, subject, html, excel_path, html_path)
    logger.info(f"周报已{'发送' if ok else '生成（SMTP未配置，跳过邮件）'}: {html_path}, {excel_path}")


if __name__ == "__main__":
    main()
