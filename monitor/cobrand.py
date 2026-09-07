"""联名内容数据包：联名看板 + 联名明细 Excel + 媒体分卷 zip，邮件发送。

用法：uv run python -m monitor.cobrand [天数，可选：仅最近N天，默认全部]
"""
import sys
from datetime import datetime
from pathlib import Path

from jinja2 import Environment, FileSystemLoader
from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .config import BASE_DIR, DATA_DIR, load_config
from .dashboard import build_dashboard
from .db import connect
from .media import _extract_media, build_zip_parts
from .mailer import _send

TEMPLATE_DIR = BASE_DIR / "monitor" / "templates"
HEADERS = ["联名", "账号", "标题", "发布时间", "点赞", "收藏", "评论", "分享", "话题", "链接"]


def _fmt_ts(ms):
    if not ms:
        return ""
    return datetime.fromtimestamp(ms / 1000).strftime("%Y-%m-%d %H:%M")


def fetch_cobrand(conn, days: int | None):
    if days:
        since_ms = int((datetime.now().timestamp() - days * 86400) * 1000)
        rows = conn.execute(
            "SELECT * FROM notes WHERE cobrand != '' AND publish_time >= ? ORDER BY publish_time DESC",
            (since_ms,)).fetchall()
    else:
        rows = conn.execute(
            "SELECT * FROM notes WHERE cobrand != '' ORDER BY publish_time DESC").fetchall()
    notes = []
    for r in rows:
        n = dict(r)
        n["title"] = n["title"] or "无标题"
        n["images"] = _extract_media(n.get("detail_json"))[0][:3]
        notes.append(n)
    return notes


def build_excel(notes, path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "联名明细"
    for col, h in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFD700")
    for n in notes:
        ws.append([n["cobrand"], n["nickname"], n["title"], _fmt_ts(n["publish_time"]),
                   n["liked_count"], n["collected_count"], n["comment_count"],
                   n["share_count"], n["tags"], n["note_url"]])
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(width + 4, 60)
    wb.save(path)
    return path


def _summary_html(notes) -> str:
    env = Environment(loader=FileSystemLoader(TEMPLATE_DIR))
    pairs = {}
    for n in notes:
        for p in n["cobrand"].split(";"):
            pairs[p.strip()] = pairs.get(p.strip(), 0) + 1
    top_pairs = sorted(pairs.items(), key=lambda x: -x[1])[:20]
    return env.get_template("cobrand_summary.html.j2").render(
        notes=notes, top_pairs=top_pairs, fmt_ts=_fmt_ts)


def main():
    days = int(sys.argv[1]) if len(sys.argv) > 1 else 0
    conn = connect()
    notes = fetch_cobrand(conn, days or None)
    conn.close()
    stamp = "全部" if not days else f"近{days}天"
    logger.info(f"联名笔记 {len(notes)} 篇（{stamp}），开始生成数据包…")

    dashboard_path = build_dashboard(cobrand_only=True)
    excel_path = build_excel(notes, DATA_DIR / "excel" / f"联名明细_{stamp}.xlsx")
    summary_html = _summary_html(notes)
    summary_path = DATA_DIR / "reports" / f"cobrand_{stamp}.html"
    summary_path.write_text(summary_html, encoding="utf-8")

    zip_paths = build_zip_parts(notes, DATA_DIR / "media_zip", f"联名媒体_{stamp}")
    logger.info(f"媒体分卷: {len(zip_paths)} 个 -> {[f'{p.stat().st_size/1048576:.0f}MB' for p in zip_paths]}")

    cfg = load_config()
    n = len(zip_paths)
    total = f"{sum(p.stat().st_size for p in zip_paths)/1048576:.0f}MB"
    for i, zp in enumerate(zip_paths, 1):
        subject = f"【XHS竞品监测】联名数据包 {stamp}（{i}/{n}，共{total}）"
        extra = [excel_path, summary_path, dashboard_path] if i == 1 else []
        ok = _send(cfg, subject, summary_html, [zp] + extra)
        logger.info(f"已发送 {i}/{n}: {zp.name} ({'ok' if ok else 'FAIL'})")
    logger.success(f"联名数据包发送完成: {n} 封邮件, 看板 {dashboard_path}")


if __name__ == "__main__":
    main()
